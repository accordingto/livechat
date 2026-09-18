#!/usr/bin/env python3
"""
把題庫試算表的題目收回成 game-data.js 的內建預設題目。

方向跟 build-sheet-template.py 相反，兩支是一組的：

    game-data.js  ──build-sheet-template.py──▶  prompt-sheet-template.xlsx
    game-data.js  ◀──sheet-to-game-data.py───   （主持人改過的試算表）

日常的題目編輯是在 Google 試算表上做的，那份才是工作中的版本；game-data.js
是「網站讀不到試算表時用的」那一份，會慢慢落後。要讓兩邊回到同一個內容就
跑這支，把試算表下載成 .xlsx 之後餵給它。

用法：
    python3 tools/sheet-to-game-data.py 下載下來的.xlsx
    python3 tools/sheet-to-game-data.py 下載下來的.xlsx --dry-run   # 只看會變什麼

分頁名稱認兩種：遊戲名（Say It Without Saying It）與舊的 GAME_DATA key
（taboo），跟 sheet-data.js 的 load() 一樣。找不到的分頁就不動那個資料集，
不會把它清空——試算表少一個分頁不該讓一款遊戲沒題目。

每一列都跑 sheet-data.js 真正的 parseTab()，不在這裡另外寫一份解析，
所以「網站讀試算表會拿到什麼」跟「寫進 game-data.js 的是什麼」必然一致。
改寫時只換掉 `key: [` 到 `],` 之間的那一段，各資料集上方的說明註解原封不動。

需要 openpyxl（只有這支維護工具需要）：pip install openpyxl
"""
import csv, io, json, os, re, subprocess, sys, tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TARGET = os.path.join(ROOT, 'game-data.js')

def node(expr, payload=None):
    """兩支 .js 串成同一個 script 再跑：game-data.js 的 `const GAME_DATA` 是
       檔案作用域的，require() 進來看不到，同一個 script 才看得到。"""
    def read(name):
        with open(os.path.join(ROOT, name), encoding='utf8') as f: return f.read()
    shim = ('global.window = {};\n'
            'global.localStorage = { getItem: () => null, setItem: () => {}, removeItem: () => {} };\n')
    src = shim + read('game-data.js') + '\n;\n' + read('sheet-data.js') + '\n;\n'
    if payload is not None:
        src += 'const PAYLOAD = ' + json.dumps(payload) + ';\n'
    src += expr
    with tempfile.NamedTemporaryFile('w', suffix='.cjs', delete=False, encoding='utf8') as f:
        f.write(src); path = f.name
    try:
        out = subprocess.run(['node', path], capture_output=True, text=True)
        if out.returncode != 0:
            raise SystemExit('node 執行失敗：\n' + (out.stderr or out.stdout))
        return json.loads(out.stdout)
    finally:
        os.unlink(path)

def sheet_specs():
    return node("console.log(JSON.stringify(window.SHEET_DATA.TABS.map("
                "k => ({ key: k, tab: window.SHEET_DATA.tabName(k) }))));")

def xlsx_to_csv(path):
    """每個分頁 → CSV 字串，交給網站真正的 CSV 解析器處理引號與換行。"""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    out = {}
    for ws in wb.worksheets:
        buf = io.StringIO(); w = csv.writer(buf, lineterminator='\n')
        empty = True
        for row in ws.iter_rows(values_only=True):
            if all(v in (None, '') for v in row): continue
            # 整數別寫成 45.0，seconds 那欄會變成字串 "45.0"
            w.writerow(['' if v is None else
                        (str(int(v)) if isinstance(v, float) and v.is_integer() else str(v))
                        for v in row])
            empty = False
        if not empty: out[ws.title] = buf.getvalue()
    return out

# ── 序列化：每個資料集照 game-data.js 本來就在用的排版 ────────────────────
def q(s):
    s = str(s).replace('\\', '\\\\')
    if "'" not in s: return "'" + s + "'"
    if '"' not in s: return '"' + s + '"'
    return '"' + s.replace('"', '\\"') + '"'

def qd(s):                                    # 固定用雙引號（conquest/hottake 的風格）
    s = str(s).replace('\\', '\\\\').replace('"', '\\"')
    return '"' + s + '"'

def grouped(items, keyfn, label):
    """照第一次出現的順序分組，每組前面插一行 /* ── 標題 ── */ 分隔。"""
    order, buckets = [], {}
    for it in items:
        k = keyfn(it)
        if k not in buckets: order.append(k); buckets[k] = []
        buckets[k].append(it)
    return [(label(k), buckets[k]) for k in order]

def ser(key, rows):
    L = []
    if key == 'hottake':
        L += ['    %s,' % qd(r) for r in rows]

    elif key == 'sophies':
        for r in rows:
            L.append('    { cat: %s,' % q(r.get('cat', '')))
            L.append('      situation: %s,' % q(r['situation']))
            L.append('      a: %s,' % q(r['a']))
            L.append('      b: %s,' % q(r['b']))
            L.append('    },')

    elif key == 'persuade':
        for r in rows:
            L.append('    { cat:%s, judge:%s, emoji:%s,'
                     % (q(r.get('cat', '')), q(r['judge']), q(r.get('emoji', ''))))
            L.append('      team:%s,' % q(r['team']))
            L.append('      situation:%s },' % q(r['situation']))

    elif key == 'scene':
        for r in rows:
            L.append('    {')
            L.append('      cat: %s,' % q(r.get('cat', '')))
            L.append('      emoji: %s,' % q(r.get('emoji', '')))
            L.append('      title: %s,' % q(r['title']))
            L.append('      situation: %s,' % q(r['situation']))
            for slot in 'abcdef':
                role = r.get(slot) or {}
                L.append('      %s: { label: %s, hint: %s },'
                         % (slot, q(role.get('label', '')), q(role.get('hint', ''))))
            L.append('    },')

    elif key == 'forbidden':
        L += ['    { word: %s, emoji: %s },' % (q(r['word']), q(r['emoji'])) for r in rows]

    elif key in ('conquest', 'conquestTruth'):
        for i, (title, items) in enumerate(
                grouped(rows, lambda r: r.get('type', ''), lambda k: k)):
            if i: L.append('')
            L.append('    /* ── %s ── */' % title)
            L += ['    { type: %s, cat: %s, text: %s, seconds: %d },'
                  % (qd(r.get('type', '')), qd(r.get('cat', '')), qd(r['text']), int(r['seconds']))
                  for r in items]

    elif key == 'taboo':
        LEVEL = {'easy': '🟢 EASY', 'medium': '🟡 MEDIUM', 'hard': '🔴 HARD'}
        by = {'easy': [], 'medium': [], 'hard': []}
        for r in rows: by[r['level']].append(r)
        for i, lv in enumerate(['easy', 'medium', 'hard']):
            if not by[lv]: continue
            if i: L.append('')
            L.append('    /* ── %s ── */' % LEVEL[lv])
            L += ['    { level: %s, emoji: %s, word: %s, forbidden: [%s] },'
                  % (q(r['level']), q(r['emoji']), q(r['word']),
                     ', '.join(q(f) for f in r['forbidden']))
                  for r in by[lv]]

    elif key == 'kangaroo':
        L += ['    { emoji: %s, charge: %s },' % (q(r['emoji']), qd(r['charge'])) for r in rows]

    elif key == 'kangarooWitness':
        L += ['    { emoji: %s, name: %s, evidence: %s },'
              % (q(r['emoji']), q(r['name']), q(r['evidence'])) for r in rows]

    else:
        raise SystemExit('不認得的資料集：' + key)
    return L

# ── 只換掉 `key: [` 到 `],` 之間那一段，註解全部留著 ──────────────────────
def splice(src, key, body):
    lines = src.split('\n')
    try:
        start = next(i for i, l in enumerate(lines) if l == '  %s: [' % key)
    except StopIteration:
        raise SystemExit('在 game-data.js 裡找不到 `%s: [`' % key)
    end = next(i for i in range(start + 1, len(lines)) if lines[i] == '  ],')
    return '\n'.join(lines[:start + 1] + body + lines[end:])

def main():
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    dry = '--dry-run' in sys.argv
    if not args:
        raise SystemExit(__doc__.strip())
    xlsx = args[0]
    if not os.path.exists(xlsx):
        raise SystemExit('找不到檔案：' + xlsx)

    specs = sheet_specs()
    tabs = xlsx_to_csv(xlsx)

    # 分頁名先認遊戲名，再認舊的 key——跟 sheet-data.js 的 load() 同一套
    picked, missing = {}, []
    for s in specs:
        name = s['tab'] if s['tab'] in tabs else (s['key'] if s['key'] in tabs else None)
        if name is None: missing.append(s); continue
        picked[s['key']] = tabs[name]

    parsed = node("""
      const S = window.SHEET_DATA, out = {};
      for (const [key, csv] of Object.entries(PAYLOAD)) {
        try { const r = S.parseTab(key, csv); out[key] = { rows: r.rows, skipped: r.skipped }; }
        catch (e) { out[key] = { error: String(e.message || e) }; }
      }
      console.log(JSON.stringify(out));
    """, picked)

    current = node("console.log(JSON.stringify(Object.fromEntries("
                   "window.SHEET_DATA.TABS.map(k => [k, (GAME_DATA[k] || []).length]))));")

    src = open(TARGET, encoding='utf8').read()
    changed, failed = [], []
    print('%-18s %7s %7s   %s' % ('資料集', '試算表', '現行', '結果'))
    print('-' * 62)
    for s in specs:
        key, cur = s['key'], current[s['key']]
        if key not in picked:
            print('%-18s %7s %7d   分頁不在檔案裡，維持原樣' % (key, '—', cur)); continue
        r = parsed[key]
        if r.get('error'):
            failed.append(key)
            print('%-18s %7s %7d   ❌ %s' % (key, '—', cur, r['error'])); continue
        rows = r['rows']
        if not rows:
            failed.append(key)
            print('%-18s %7d %7d   ❌ 一列都沒有，維持原樣' % (key, 0, cur)); continue
        src = splice(src, key, ser(key, rows))
        changed.append(key)
        note = '寫入 (%+d)' % (len(rows) - cur)
        if r['skipped']: note += '，跳過 %d 列（必填欄位空白）' % r['skipped']
        print('%-18s %7d %7d   %s' % (key, len(rows), cur, note))
    print('-' * 62)

    for s in missing:
        print('（試算表沒有「%s」這個分頁）' % s['tab'])

    if dry:
        print('\n--dry-run：沒有寫檔'); return 0
    if not changed:
        print('\n沒有任何資料集可以寫入'); return 1
    open(TARGET, 'w', encoding='utf8').write(src)
    print('\n已寫入 %s（%d 個資料集）' % (os.path.relpath(TARGET, ROOT), len(changed)))
    return 1 if failed else 0

if __name__ == '__main__':
    sys.exit(main())
