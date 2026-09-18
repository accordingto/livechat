#!/usr/bin/env python3
"""
從 game-data.js ＋ sheet-data.js 的 SCHEMA 產生 prompt-sheet-template.xlsx。

這份範本的內容就是「網站沒讀取 Google Sheet 時用的預設題目」，所以它不是手寫
維護的，是每次題庫改動之後重跑這支程式產生出來的——CLAUDE.md 要求「改題庫時
這個檔案要跟著重新產生」，這支程式就是那個步驟。

分頁名稱、欄位、順序全部讀 sheet-data.js 的 SCHEMA，不在這裡重寫一份，
兩邊才不會漂。產生完會自動做一次往返驗證：
xlsx → CSV → sheet-data.js 真正的 parseTab() → 跟 GAME_DATA 逐欄比對。

用法：
    python3 tools/build-sheet-template.py            # 產生並驗證
    python3 tools/build-sheet-template.py --verify   # 只驗證現有的檔案

需要 openpyxl（只有這支維護工具需要，網站本身沒有任何相依）：
    pip install openpyxl
"""
import json, subprocess, sys, tempfile, os

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT  = os.path.join(ROOT, 'prompt-sheet-template.xlsx')

# ── 從瀏覽器用的那兩支 .js 直接取值，不另外維護一份副本 ────────────────────
SHIM = """
global.window = {};
global.localStorage = { getItem: () => null, setItem: () => {}, removeItem: () => {} };
"""

def node(expr):
    """兩支 .js 串成同一個 script 再跑：game-data.js 的 `const GAME_DATA` 是
       檔案作用域的，require() 進來看不到，同一個 script 才看得到。"""
    def read(name):
        with open(os.path.join(ROOT, name), encoding='utf8') as f: return f.read()
    src = SHIM + read('game-data.js') + '\n' + read('sheet-data.js') + '\n' + expr
    with tempfile.NamedTemporaryFile('w', suffix='.cjs', delete=False, encoding='utf8') as f:
        f.write(src); path = f.name
    try:
        out = subprocess.run(['node', path], capture_output=True, text=True)
        if out.returncode != 0:
            raise SystemExit('node 執行失敗：\n' + (out.stderr or out.stdout))
        return json.loads(out.stdout)
    finally:
        os.unlink(path)

def load_spec():
    return node("""
      const S = window.SHEET_DATA;
      console.log(JSON.stringify({
        tabs: S.TABS.map(k => ({
          key: k, tab: S.tabName(k),
          cols: S.SCHEMA[k].cols, required: S.SCHEMA[k].required,
          count: (GAME_DATA[k] || []).length,
        })),
        data: Object.fromEntries(S.TABS.map(k => [k, GAME_DATA[k] || []])),
      }));
    """)

# ── GAME_DATA 的物件 → 試算表的一列 ──────────────────────────────────────
def to_row(key, cols, item):
    if key == 'hottake':                       # 純字串陣列，不是物件
        return [item]
    row = []
    for c in cols:
        if key == 'taboo' and c.startswith('forbidden'):
            i = int(c[len('forbidden'):]) - 1
            f = item.get('forbidden') or []
            row.append(f[i] if i < len(f) else '')
        elif key == 'scene' and ('_label' in c or '_hint' in c):
            slot, part = c.split('_')
            role = item.get(slot) or {}
            row.append(role.get(part, ''))
        else:
            row.append(item.get(c, ''))
    return row

# ── README 分頁 ─────────────────────────────────────────────────────────
def readme_lines(tabs):
    L = [
        ('h1', 'IceBreak Hub 題庫'),
        ('', ''),
        ('p',  '網站七款遊戲的題目（💬 Let\'s Talk 的話題不走這份試算表）。這裡的內容'),
        ('p',  '就是網站沒有讀取試算表時所使用的內建預設題目——你是在既有的題目上'),
        ('p',  '修改，不是從零開始填。'),
        ('', ''),
        ('h2', '怎麼讓網站讀到'),
        ('li', '1. 上傳到 Google Drive，用「Google 試算表」開啟（已經在試算表裡就跳過）。'),
        ('li', '2. 右上角「共用」→ 一般存取權改成「知道連結的任何人」→ 權限「檢視者」。'),
        ('warn', '     ⚠️ 少了這一步，網站讀不到，只會安靜地退回內建題庫。'),
        ('li', '3. 回網站首頁 →「設定房間」走到最後一步 → 按「📄 讀取線上題庫」。'),
        ('', ''),
        ('h2', '改題目前要知道的五件事'),
        ('li', '• 分頁名稱＝遊戲名稱，不要改，網站是靠名稱找資料的。'),
        ('li', '• 第一列的標題不要改；黃色底的是必填欄位，該欄留空的那一列會被跳過。'),
        ('li', '• 欄位順序可以調換，網站看的是標題文字而不是位置。'),
        ('li', '• 題目可以自由增減，不必維持原本的題數；整個分頁清空＝那款用內建題庫。'),
        ('li', '• 只有按下「讀取」的當下才會連線；讀到的內容存在那一台裝置上，'),
        ('li', '  換一台電腦要各自再讀一次。'),
        ('', ''),
        ('h2', '寫題目的原則'),
        ('p',  '玩家是在語音聊天室裡練英文口說的初學～中階學習者。題目是要被唸出來、'),
        ('p',  '當場聽懂、當場回答的，所以句子要短、用字要常見，避免美式俚語或需要'),
        ('p',  '文化背景才看得懂的說法。參與者沒有鏡頭，一律寫成「唸出來／說出來」，'),
        ('p',  '不要寫 show us。'),
        ('', ''),
        ('h2', '各分頁'),
    ]
    note = {
        'taboo': 'word 是目標字，forbidden1～6 是禁字，level 只能填 easy／medium／hard',
        'hottake': '一句故意講得很極端的意見，讓大家選邊站',
        'sophies': 'situation 是兩難情境，a／b 是兩個選項',
        'persuade': 'judge 是裁判的身分，team 是說服方的身分',
        'scene': 'a_label～f_label 是六個角色名，a_hint～f_hint 是演法提示',
        'conquest': 'type 是五大類型，cat 是這題自己的小標題，seconds 一律 45',
        'conquestTruth': '欄位與大冒險相同',
        'kangaroo': 'charge 會接在「The defendant is charged with…」後面',
        'kangarooWitness': 'name 是證人角色名，evidence 是他帶來的一句證詞',
        'forbidden': '不是遊戲：發完連結後用來核對誰拿到哪張卡的日常單字',
    }
    for t in tabs:
        L.append(('row', (t['tab'], note.get(t['key'], ''), f"{t['count']} 題")))
    return L

def build(spec):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    FONT = 'Arial'
    wb = Workbook()
    ws = wb.active
    ws.title = '說明 README'

    head  = Font(name=FONT, bold=True, size=15, color='1A1A38')
    sub   = Font(name=FONT, bold=True, size=12, color='C0392B')
    body  = Font(name=FONT, size=11)
    warn  = Font(name=FONT, size=11, color='B7791F')
    tabf  = Font(name=FONT, size=11, bold=True)

    for i, (kind, val) in enumerate(readme_lines(spec['tabs']), start=1):
        if kind == 'row':
            for j, v in enumerate(val, start=1):
                c = ws.cell(row=i, column=j, value=v)
                c.font = tabf if j == 1 else body
                c.alignment = Alignment(vertical='top')
        else:
            c = ws.cell(row=i, column=1, value=val)
            c.font = {'h1': head, 'h2': sub, 'warn': warn}.get(kind, body)
    ws.column_dimensions['A'].width = 62
    ws.column_dimensions['B'].width = 58
    ws.column_dimensions['C'].width = 10
    ws.sheet_view.showGridLines = False

    hdr_font = Font(name=FONT, bold=True, size=11, color='1A1A38')
    req_fill = PatternFill('solid', fgColor='FFE9A8')     # 必填
    opt_fill = PatternFill('solid', fgColor='E8E8F0')     # 選填
    WIDE = {'situation', 'text', 'charge', 'evidence', 'a', 'b', 'team', 'judge'}
    MID  = {'title', 'cat', 'type', 'word', 'name'}

    for t in spec['tabs']:
        ws = wb.create_sheet(t['tab'])
        req = set(t['required'])
        for j, col in enumerate(t['cols'], start=1):
            c = ws.cell(row=1, column=j, value=col)
            c.font = hdr_font
            c.fill = req_fill if col in req else opt_fill
            c.alignment = Alignment(vertical='center')
            w = 70 if col in WIDE else 22 if col in MID else 16
            if col.endswith('_hint'):  w = 34
            if col.endswith('_label'): w = 22
            if col == 'emoji' or col.startswith('emoji'): w = 8
            if col == 'seconds': w = 10
            ws.column_dimensions[get_column_letter(j)].width = w

        for item in spec['data'][t['key']]:
            ws.append(to_row(t['key'], t['cols'], item))
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = body
                c.alignment = Alignment(vertical='top', wrap_text=True)
        ws.freeze_panes = 'A2'

    wb.save(OUT)
    return OUT

# ── 往返驗證：xlsx → CSV → 真正的 parseTab() → 跟 GAME_DATA 比對 ──────────
def verify(spec):
    import csv, io
    from openpyxl import load_workbook
    wb = load_workbook(OUT)
    sheets = {}
    for t in spec['tabs']:
        ws = wb[t['tab']]
        buf = io.StringIO()
        w = csv.writer(buf, lineterminator='\n')
        for row in ws.iter_rows(values_only=True):
            if all(v in (None, '') for v in row): continue
            w.writerow(['' if v is None else str(v) for v in row])
        sheets[t['key']] = buf.getvalue()

    res = node("""
      const sheets = %s;
      const S = window.SHEET_DATA;
      const out = {};
      for (const key of S.TABS) {
        let parsed, err = null;
        try { parsed = S.parseTab(key, sheets[key]); }
        catch (e) { err = String(e.message || e); parsed = { rows: [], skipped: 0 }; }
        out[key] = {
          error: err, skipped: parsed.skipped,
          match: JSON.stringify(parsed.rows) === JSON.stringify(GAME_DATA[key] || []),
          got: parsed.rows.length, want: (GAME_DATA[key] || []).length,
        };
      }
      console.log(JSON.stringify(out));
    """ % json.dumps(sheets))

    ok = True
    print(f'\n往返驗證  xlsx → CSV → parseTab() → GAME_DATA')
    print('─' * 62)
    for t in spec['tabs']:
        r = res[t['key']]
        good = r['match'] and not r['error'] and not r['skipped']
        ok &= good
        mark = '✅' if good else '❌'
        extra = f"  {r['error']}" if r['error'] else (f"  跳過 {r['skipped']} 列" if r['skipped'] else '')
        print(f"{mark} {t['tab']:<30} {r['got']:>4} / {r['want']:<4} 題{extra}")
    print('─' * 62)
    print('全部逐字元相符' if ok else '有分頁對不上，上面標 ❌ 的要查')
    return ok

if __name__ == '__main__':
    spec = load_spec()
    if '--verify' not in sys.argv:
        print('產生', build(spec))
    sys.exit(0 if verify(spec) else 1)
